"""Shared Word formatting helpers for academic manuscripts."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, east_asia=None):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    ea = east_asia or name
    rFonts.set(qn("w:eastAsia"), ea)
    rFonts.set(qn("w:cs"), name)


def pf(p, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, first=None, line=1.15, left=0):
    f = p.paragraph_format
    f.alignment = align
    f.space_before = Pt(before)
    f.space_after = Pt(after)
    f.line_spacing = line
    f.left_indent = Inches(left)
    if first is not None:
        f.first_line_indent = Inches(first)


def add(p, text, **kw):
    r = p.add_run(text)
    set_run_font(r, **kw)
    return r


def setup_doc(chinese=False):
    doc = Document()
    s = doc.sections[0]
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1.0)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun" if chinese else "Times New Roman")
    return doc


def H(doc, text, level=1, ea="Times New Roman"):
    sizes = {1: 14, 2: 12, 3: 12}
    p = doc.add_paragraph()
    pf(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=12 if level == 1 else 8, after=6, first=0)
    add(p, text, size=sizes.get(level, 12), bold=True, italic=(level >= 3), east_asia=ea)
    return p


def P(doc, text, ea="Times New Roman", first=0.3, size=12):
    p = doc.add_paragraph()
    pf(p, first=first)
    add(p, text, size=size, east_asia=ea)
    return p


def Pb(doc, lead, text, ea="Times New Roman", first=0.3):
    p = doc.add_paragraph()
    pf(p, first=first)
    add(p, lead, bold=True, east_asia=ea)
    add(p, text, east_asia=ea)
    return p


def cap(doc, text, ea="Times New Roman"):
    p = doc.add_paragraph()
    pf(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=10, first=0)
    add(p, text, size=10, italic=True, east_asia=ea)


def fig(doc, path: Path, width=5.6):
    p = doc.add_paragraph()
    pf(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=2, first=0)
    p.add_run().add_picture(str(path), width=Inches(width))


def table(doc, headers, rows, ea="Times New Roman"):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        pf(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=1, after=1, first=0)
        add(p, h, size=9, bold=True, east_asia=ea)
    for i, row in enumerate(rows):
        hi = any(x in str(row[0]) for x in ("Ours", "本文", "Meeting"))
        for j, val in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            pf(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=1, after=1, first=0)
            add(p, str(val), size=9, bold=hi, east_asia=ea)
    return t


def algo(doc, title, lines, ea="Times New Roman"):
    p = doc.add_paragraph()
    pf(p, first=0, before=6, after=3)
    add(p, title, bold=True, size=10, east_asia=ea)
    for line in lines:
        p = doc.add_paragraph()
        pf(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=0, first=0, line=1.0, left=0.2)
        add(p, line, size=9, name="Courier New", east_asia="Courier New")


def refs_block(doc, refs, title="References", ea="Times New Roman"):
    H(doc, title, 1, ea)
    for i, r in enumerate(refs, 1):
        p = doc.add_paragraph()
        pf(p, before=0, after=3, first=0, line=1.0, left=0.35)
        p.paragraph_format.first_line_indent = Inches(-0.35)
        add(p, f"[{i}] {r}", size=9, east_asia=ea)
