"""Fill CISC7298 Progress Report 1 using the official template layout.

The school template uses one 1x1 content table under each section heading.
We keep that form, write section text into those tables, and insert a
results summary table inside Experimental Results. No References section.
"""
from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_TEMPLATE = Path(
    r"C:\Users\zkb17\xwechat_files\wxid_g3r567e59c9a22_605c\msg\file\2026-09"
    r"\CISC7298 Progress Report 1(1).docx"
)
OUT_DOCS = ROOT / "docs" / "CISC7298_Progress_Report_1.docx"
OUT_ROOT = ROOT / "CISC7298_Progress_Report_1.docx"

SECTION_CONTENTS = {
    "Title and Abstract:": [
        (
            "Title: Hierarchical Lifelong Multi-Agent Pickup-and-Delivery for "
            "Warehouse AGVs via Spatiotemporal A* with ECBS Escalation"
        ),
        (
            "Abstract. This project studies lifelong multi-agent pickup-and-delivery "
            "(MAPD) for automated guided vehicles (AGVs) on a 20x20 warehouse grid. "
            "Station queues follow FIFO order. Turns are performed in place and finish "
            "within one second (including 180 degrees), without moving in the same second. "
            "Pickup and unload each require a one-second dwell. We propose a "
            "hierarchical planner: a fast spatiotemporal A* layer (M0) runs by "
            "default, and windowed Enhanced Conflict-Based Search (ECBS) is used only "
            "when progress stalls. Two small neural networks support the hierarchy: "
            "SwapNet for station task swaps under M0, and StallEscalateNet for "
            "M0-to-ECBS handoff. All trajectories are checked by a shared validator. "
            "On twenty maps with 100 tasks each, both hierarchical M0-ECBS and pure "
            "windowed ECBS finish all tasks with valid trajectories. Reactive PIBT, "
            "M0 alone, and priority planning do worse on congested maps. Hierarchy "
            "also lowers average makespan versus always-on ECBS (about 4616 vs 7364 "
            "ticks), at a moderate increase in planning time."
        ),
    ],
    "Introduction (with OR w/o Related Work/Literature Review):": [
        (
            "Warehouse AGVs must keep assigning pickups, avoid static obstacles and "
            "other robots, and respect FIFO station order. One-shot multi-agent path "
            "finding (MAPF) does not cover lifelong task release. Reactive methods "
            "such as PIBT scale well but often deadlock in narrow corridors. Joint "
            "search methods such as CBS and ECBS resolve conflicts more reliably, "
            "but calling them at every step can raise makespan and runtime."
        ),
        (
            "Our goal is to complete as many deliveries as possible under these "
            "motion and FIFO rules. The main question is whether a hierarchical "
            "policy can match always-ECBS completion: use cheap spatiotemporal A* "
            "most of the time, and call ECBS only when throughput stalls, so that "
            "easy maps still keep M0-level efficiency. Two small neural aides help "
            "this policy without replacing search."
        ),
        (
            "Related work. Conflict-Based Search (CBS) and weighted ECBS are standard "
            "for bounded-suboptimal joint plans. Priority planning and windowed "
            "lifelong planners (WHCA*/RHCR-style) support continuous operation with "
            "limited horizons. PIBT offers fast one-step conflict resolution without "
            "reservations. Learned components in MAPF are often large policies; we "
            "instead use narrow binary networks as gates inside a classical stack. "
            "M0 runs lifelong traffic with reservation tables. StallEscalateNet and "
            "rule stall signals decide when to start ECBS. A handoff snapshot keeps "
            "cargo state and FIFO heads consistent across the switch."
        ),
    ],
    "Methodology:": [
        (
            "Problem model. Each map has six pickup stations and sixteen dropoff "
            "hubs on a 20x20 grid. The main study uses 100-task instances; larger "
            "400-task runs are used for stress tests. An AGV carries at most one "
            "package. Pickup happens at the approach cell next to a station after a "
            "one-second dwell with no turning. Unload happens on a free 4-neighbor "
            "of the destination hub, with the same dwell rule."
        ),
        (
            "Layer M0 (default). Free AGVs are assigned to station FIFO heads, with "
            "at most one approacher per station. Paths are planned by spatiotemporal "
            "A* with a rolling moving-obstacle horizon. Vertex collisions and edge "
            "swaps are forbidden. M0 is fast on open maps, but it may stall when "
            "corridors become congested."
        ),
        (
            "Hierarchical gate and ECBS escalation. The system starts in M0 and "
            "tracks runtime signals such as completion plateaus and assignment "
            "failures. If M0 still stalls after a short recovery window, or if "
            "StallEscalateNet fires, we take a handoff snapshot and continue with "
            "windowed joint ECBS. FIFO queue order is never changed by escalation."
        ),
        (
            "AI module 1: SwapNet. SwapNet is a small MLP with structure "
            "16 -> 48 -> 48 -> 1 and a sigmoid output. Input is a 16-dimensional "
            "feature vector for a candidate station swap under M0: normalized ETAs "
            "and distances of the far and near agents, their ETA gap, task urgency, "
            "remaining time and SLA slack, urgent gain, whether the near agent is "
            "idle or already holds an urgent task, queue depth, margin ratio, and "
            "SLA meet/miss flags. Output is P(authorize swap). Training imitates a "
            "heuristic teacher on logged swap contexts and favors precision over "
            "recall. At runtime the engine still rejects any swap that would skip a "
            "FIFO station head, so the network cannot break queue legality."
        ),
        (
            "AI module 2: StallEscalateNet. StallEscalateNet is an MLP with structure "
            "15 -> 48 -> 48 -> 1 and a sigmoid output. Input is a 15-dimensional "
            "stall vector: map hardness and narrow-cut priors, done ratio, "
            "normalized plateau / idle / assign-fail / throughput-drop / wall-tick "
            "streaks, plan-fail rate, wall-window hit, sim-per-task, wall elapsed, "
            "has-work, loaded ratio, and recent completions in a short sim window. "
            "Output is P(escalate M0 to ECBS), compared with a high threshold "
            "(about 0.72) so that healthy M0 on easy maps is left alone. The "
            "network is trained to catch wall-clock thrash while avoiding unnecessary "
            "escalation."
        ),
        (
            "Other AI modules. WaveNet (MLP on 12-D agent features, optional map CNN) "
            "and SceneDifficultyNet (map CNN plus scalars) exist for wave ranking and "
            "scene hardness estimation. They are offline or optional in this report "
            "and are not required for the hierarchical baseline numbers below. "
            "Search itself remains classical: M0 spatiotemporal A* and windowed ECBS."
        ),
        (
            "Windowed ECBS baseline. We also run pure windowed ECBS as a joint "
            "planning baseline. To keep trajectories valid, we added a few practical "
            "guards: reject unload events off the legal pad; restore cargo state when "
            "the log still shows loaded but the planner bag lost the task; and refuse "
            "a new pickup while another task is still carried. These fixes remove "
            "premature unload and FIFO errors seen on dense maps such as SH11 and "
            "SH12."
        ),
        (
            "Baselines. Under the same maps, task files, motion rules, and validator, "
            "we compare five methods: (1) reactive PIBT; (2) M0-only spatiotemporal "
            "A*; (3) priority planning with A*; (4) windowed ECBS without hierarchy; "
            "(5) hierarchical M0-ECBS with SwapNet and StallEscalateNet (ours)."
        ),
    ],
    "Experimental Results:": [
        (
            "Setup. We evaluate maps SH_custom_01-20 with 100-task instances and the "
            "full AGV fleet. Constrained runs use a 1800 s wall-time limit; "
            "hierarchical and ECBS runs that need more planning time are repeated "
            "without that limit. Simulation horizons go up to 500000 ticks when we "
            "measure completion. Metrics are completed tasks out of 100, trajectory "
            "validity, simulation makespan, wall-clock time, and total Manhattan "
            "travel."
        ),
        (
            "Aggregate results on SH01-20 are shown below. Both ECBS and Hier finish "
            "all 100 tasks on every map with valid trajectories. Hierarchy cuts "
            "average makespan by about 37% versus always-on ECBS (4616 vs 7364), "
            "with a moderate rise in average wall-clock time (704 s vs 543 s). PIBT "
            "is very fast but averages only 63.9 completed tasks and fully finishes "
            "8/20 maps. Pure M0 stays valid but finishes only 4/20 maps. Priority "
            "planning is weakest on both completion and validity (7/20 valid)."
        ),
        # marker for results table insertion
        "__RESULTS_TABLE__",
        (
            "Map-level notes. On easy SH01, M0 and Hier finish with low makespan "
            "(about 546-612 ticks), while PIBT and ECBS need about 2200 ticks. "
            "Priority planning finishes 100 tasks but fails FIFO checks. On narrow "
            "SH03, PIBT, M0, and priority planning complete at most five tasks, "
            "while ECBS and Hier both reach 100/100 and stay valid. On SH04, Hier "
            "stays close to M0 (670 vs 684) instead of paying the longer ECBS plan "
            "(4219). On dense SH11 and SH12, after the validity fixes, ECBS reaches "
            "100/100 (makespan 18021 and 16981), and Hier matches completion with "
            "shorter makespan (4775 and 5220)."
        ),
        (
            "Discussion. Reactive and local search lose throughput under congestion. "
            "Always-on ECBS is reliable but costly in simulation time. Hierarchical "
            "M0-first planning with StallEscalateNet handoff matches ECBS on "
            "completion and validity, while recovering much of M0's makespan on "
            "easier maps. SwapNet mainly helps M0 serve urgent tasks without an "
            "early ECBS call; it does not replace joint search on narrow maps. "
            "This is the main empirical finding of Progress Report 1."
        ),
    ],
    "Conclusion:": [
        (
            "So far we have built a hierarchical lifelong MAPD system with two "
            "online neural aides (SwapNet and StallEscalateNet), a windowed ECBS "
            "baseline with stronger validity checks, and a full SH01-20 comparison "
            "with 100 tasks per map. Gated M0-ECBS keeps 20/20 full completions with "
            "valid trajectories and reduces average makespan relative to always-on "
            "ECBS. For Progress Report 2 we plan to: (i) report aligned 400-task "
            "stress results on the same maps; (ii) ablate SwapNet and "
            "StallEscalateNet more carefully; (iii) relate handoff timing to map "
            "structure; and (iv) study remaining high wall-time cases."
        ),
    ],
}

RESULTS_HEADERS = [
    "Method",
    "maps",
    "full 100",
    "VALID",
    "avg done",
    "avg sim",
    "avg wall (s)",
]
RESULTS_ROWS = [
    ["PIBT", "20", "8/20", "20/20", "63.9", "62811", "40.5"],
    ["M0", "20", "4/20", "20/20", "44.6", "942", "1478.5"],
    ["PP", "20", "1/20", "7/20", "32.4", "848", "1716.1"],
    ["ECBS", "20", "20/20", "20/20", "100.0", "7364", "542.7"],
    ["Hier (ours)", "20", "20/20", "20/20", "100.0", "4616", "703.7"],
]


def set_run_font(run, size=11, bold=False, name="Times New Roman"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)


def _clear_cell(cell) -> None:
    # Reset to a single empty paragraph (Word cells need ≥1 p).
    cell.text = ""
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _add_para_in_cell(cell, text: str, *, first=False, bold=False, size=11):
    if first:
        p = cell.paragraphs[0]
    else:
        p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.first_line_indent = Inches(0.25)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def _add_results_table_in_cell(cell) -> None:
    # Nested table inside the section content table.
    table = cell.add_table(rows=1 + len(RESULTS_ROWS), cols=len(RESULTS_HEADERS))
    table.style = "Table Grid"
    for j, h in enumerate(RESULTS_HEADERS):
        c = table.rows[0].cells[j]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        set_run_font(run, size=9, bold=True)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, row in enumerate(RESULTS_ROWS):
        for j, val in enumerate(row):
            c = table.rows[i + 1].cells[j]
            c.text = ""
            run = c.paragraphs[0].add_run(str(val))
            set_run_font(run, size=9)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _heading_to_table_index(doc: Document) -> dict[str, int]:
    """Map section heading text -> following table index (template order)."""
    body = doc.element.body
    headings: list[str] = []
    table_i = -1
    mapping: dict[str, int] = {}
    last_heading = None
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            texts = [t.text for t in child.iter(qn("w:t")) if t.text]
            text = "".join(texts).strip()
            if text.endswith(":") and text in SECTION_CONTENTS:
                last_heading = text
        elif tag == "tbl":
            table_i += 1
            if last_heading is not None and last_heading not in mapping:
                mapping[last_heading] = table_i
                last_heading = None
    return mapping


def fill_template(template: Path) -> Document:
    doc = Document(str(template))
    mapping = _heading_to_table_index(doc)
    if len(mapping) != 5:
        raise RuntimeError(f"Expected 5 section tables, got {mapping}")

    for heading, texts in SECTION_CONTENTS.items():
        ti = mapping[heading]
        cell = doc.tables[ti].cell(0, 0)
        _clear_cell(cell)
        first = True
        for item in texts:
            if item == "__RESULTS_TABLE__":
                _add_results_table_in_cell(cell)
                # spacer paragraph after nested table
                p = cell.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                first = False
                continue
            _add_para_in_cell(cell, item, first=first)
            first = False
    return doc


def _save(doc: Document, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(path))
        return path
    except PermissionError:
        alt = path.with_name(path.stem + "_new" + path.suffix)
        doc.save(str(alt))
        print(f"Permission denied for {path}; wrote {alt}")
        return alt


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", type=str, default=str(DEFAULT_TEMPLATE))
    ap.add_argument(
        "--also",
        action="append",
        default=[],
        help="Extra output path(s); may be repeated",
    )
    args = ap.parse_args()

    template = Path(args.template)
    if not template.exists():
        raise SystemExit(f"Template not found: {template}")

    doc = fill_template(template)
    wrote = _save(doc, OUT_DOCS)
    print(f"Wrote {wrote} ({wrote.stat().st_size} bytes)")
    try:
        shutil.copy2(wrote, OUT_ROOT)
        print(f"Copied {OUT_ROOT}")
    except PermissionError:
        alt = OUT_ROOT.with_name(OUT_ROOT.stem + "_new" + OUT_ROOT.suffix)
        shutil.copy2(wrote, alt)
        print(f"Permission denied for {OUT_ROOT}; copied {alt}")

    also = list(args.also) or [
        str(
            Path(r"C:\Users\zkb17\xwechat_files\wxid_g3r567e59c9a22_605c\msg\file\2026-09")
            / "CISC7298 Progress Report 1.docx"
        ),
        str(
            Path(r"C:\Users\zkb17\xwechat_files\wxid_g3r567e59c9a22_605c\msg\file\2026-09")
            / "CISC7298 Progress Report 1(1).docx"
        ),
    ]
    for dest_s in also:
        dest = Path(dest_s)
        try:
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(wrote, dest)
            print(f"Copied {dest}")
        except PermissionError:
            alt = dest.with_name(dest.stem + "_filled" + dest.suffix)
            shutil.copy2(wrote, alt)
            print(f"Permission denied for {dest}; copied {alt}")


if __name__ == "__main__":
    main()
